"""
make_comparison_table.py
========================
Reads results/Comparison_tables.xlsx and produces:
  1. LaTeX table  — results/table_comparison.tex
  2. XLSX table   — results/table_comparison.xlsx
  3. PNG table    — results/table_comparison.png
  4. Heatmap      — results/heatmap_comparison.png
  5. Delta heatmap— results/heatmap_delta_acc.png

Usage (from any directory):
    python scripts/make_comparison_table.py
"""

import os
import sys
import warnings
warnings.filterwarnings("ignore")

import pandas as pd
import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.colors import LinearSegmentedColormap
import openpyxl
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              numbers as xl_numbers)
from openpyxl.utils import get_column_letter

# ── Paths ──────────────────────────────────────────────────────────────────────
ROOT_DIR  = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
XLSX_IN   = os.path.join(ROOT_DIR, "results", "Comparison_tables.xlsx")
OUT_DIR   = os.path.join(ROOT_DIR, "results")

# ── Ordering & labels ──────────────────────────────────────────────────────────
METHOD_ORDER = ["Mean", "Zero", "EM", "DK_Mean", "DK_Zero", "DK_EM", "Proposed"]
METHOD_LABELS = {
    "Mean":     "Mean",
    "Zero":     "Zero",
    "EM":       "EM",
    "DK_Mean":  "DK+Mean",
    "DK_Zero":  "DK+Zero",
    "DK_EM":    "DK+EM",
    "Proposed": "Proposed",
}

DATASET_ORDER = ["iris", "alcoholqcm", "seeds", "wine", "vehicle", "glass"]
DATASET_LABELS = {
    "iris":       "Iris",
    "alcoholqcm": "AlcoholQCM",
    "seeds":      "Seeds",
    "wine":       "Wine",
    "vehicle":    "Vehicle",
    "glass":      "Glass",
}

METRICS = [
    ("ACC_mean", "ACC_sd",  "ACC (%)"),
    ("NMI_mean", "NMI_sd",  "NMI (%)"),
    ("F_mean",   "F_sd",    "F-Score (%)"),
    ("PUR_mean", "PUR_sd",  "PUR (%)"),
]

# ── Load & merge all sheets ────────────────────────────────────────────────────
def load_data(path: str) -> pd.DataFrame:
    xl = pd.ExcelFile(path)
    frames = [xl.parse(s) for s in xl.sheet_names]
    data = pd.concat(frames, ignore_index=True)
    data.columns = [c.strip() for c in data.columns]
    data["dataset"] = data["dataset"].str.lower().str.strip()
    data["method"]  = data["method"].str.strip()
    return data

def fmt(mean: float, sd: float) -> str:
    return f"{mean:.1f} ± {sd:.1f}"

def best_per_row(data: pd.DataFrame, mean_col: str) -> dict:
    best = {}
    for ds in DATASET_ORDER:
        sub = data[data["dataset"] == ds]
        if sub.empty:
            continue
        best[ds] = sub.loc[sub[mean_col].idxmax(), "method"]
    return best

def active_datasets(data: pd.DataFrame) -> list:
    return [ds for ds in DATASET_ORDER if ds in data["dataset"].values]


# ════════════════════════════════════════════════════════════════════════════════
# 1. LaTeX
# ════════════════════════════════════════════════════════════════════════════════
def make_latex(data: pd.DataFrame, out_path: str) -> None:
    col_headers = [METHOD_LABELS[m] for m in METHOD_ORDER]
    ncols = 1 + len(METHOD_ORDER)
    lines = [
        r"\begin{table*}[t]", r"\centering",
        r"\caption{Clustering performance (mean $\pm$ std, averaged over missing ratios 10\%--70\%).}",
        r"\label{tab:comparison}", r"\small",
        r"\begin{tabular}{l" + "c" * len(METHOD_ORDER) + r"}",
        r"\toprule",
        r"\textsc{Dataset} & " + " & ".join(r"\textsc{" + h + r"}" for h in col_headers) + r" \\",
    ]
    for mean_col, sd_col, label in METRICS:
        best = best_per_row(data, mean_col)
        lines += [r"\midrule",
                  r"\multicolumn{" + str(ncols) + r"}{c}{\textit{" + label + r"}} \\",
                  r"\midrule"]
        for ds in active_datasets(data):
            sub   = data[data["dataset"] == ds]
            cells = [r"\textsc{" + DATASET_LABELS[ds] + r"}"]
            for m in METHOD_ORDER:
                r_ = sub[sub["method"] == m]
                val = fmt(float(r_[mean_col].iloc[0]), float(r_[sd_col].iloc[0])) if not r_.empty else "—"
                if best.get(ds) == m:
                    val = r"\textbf{" + val + r"}"
                cells.append(val)
            lines.append(" & ".join(cells) + r" \\")
    lines += [r"\bottomrule", r"\end{tabular}", r"\end{table*}"]
    with open(out_path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print(f"  LaTeX : {out_path}")


# ════════════════════════════════════════════════════════════════════════════════
# 2. XLSX — styled, one sheet per metric + summary sheet
# ════════════════════════════════════════════════════════════════════════════════
def _xl_fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color.lstrip("#"))

def _xl_border(style="thin") -> Border:
    s = Side(style=style)
    return Border(left=s, right=s, top=s, bottom=s)

def _xl_thick_border() -> Border:
    thin  = Side(style="thin")
    thick = Side(style="medium")
    return Border(left=thin, right=thin, top=thick, bottom=thick)

# Color palette
C_HEADER_BG   = "2C3E50"
C_HEADER_FG   = "FFFFFF"
C_METRIC_BG   = "D5E8F5"
C_METRIC_FG   = "1A3A5C"
C_BEST_BG     = "FFD580"
C_BEST_FG     = "7B3F00"
C_PROPOSED_BG = "EAD7F5"
C_ODD_BG      = "FFFFFF"
C_EVEN_BG     = "F4F8FC"

def make_xlsx(data: pd.DataFrame, out_path: str) -> None:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)          # remove default blank sheet
    datasets = active_datasets(data)
    col_labels = [METHOD_LABELS[m] for m in METHOD_ORDER]

    # ── Helper: write one metric block onto a sheet ───────────────────────────
    def write_metric_block(ws, start_row: int,
                           mean_col: str, sd_col: str, metric_label: str):
        best = best_per_row(data, mean_col)
        n_cols = 1 + len(METHOD_ORDER)

        # Metric section header (merged)
        ws.merge_cells(start_row=start_row, start_column=1,
                       end_row=start_row, end_column=n_cols)
        cell = ws.cell(start_row, 1, metric_label)
        cell.font      = Font(bold=True, color=C_METRIC_FG, size=10)
        cell.fill      = _xl_fill(C_METRIC_BG)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border    = _xl_thick_border()
        ws.row_dimensions[start_row].height = 16

        row = start_row + 1
        for di, ds in enumerate(datasets):
            sub   = data[data["dataset"] == ds]
            bg    = C_ODD_BG if di % 2 == 0 else C_EVEN_BG

            # Dataset name
            c = ws.cell(row, 1, DATASET_LABELS[ds])
            c.font      = Font(italic=True, size=9)
            c.fill      = _xl_fill(bg)
            c.alignment = Alignment(horizontal="left", vertical="center",
                                    indent=1)
            c.border    = _xl_border()

            for j, m in enumerate(METHOD_ORDER, start=2):
                r_ = sub[sub["method"] == m]
                if r_.empty:
                    val_str = "—"
                    mean_val = None
                else:
                    mean_val = float(r_[mean_col].iloc[0])
                    sd_val   = float(r_[sd_col].iloc[0])
                    val_str  = fmt(mean_val, sd_val)

                c = ws.cell(row, j, val_str)
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.border    = _xl_border()

                if best.get(ds) == m:
                    c.fill = _xl_fill(C_BEST_BG)
                    c.font = Font(bold=True, color=C_BEST_FG, size=9)
                else:
                    c.fill = _xl_fill(bg)
                    c.font = Font(size=9)

            ws.row_dimensions[row].height = 15
            row += 1

        return row  # next available row

    # ── Sheet: one per metric ─────────────────────────────────────────────────
    for mean_col, sd_col, metric_label in METRICS:
        metric_name = metric_label.split()[0]          # "ACC", "NMI", etc.
        ws = wb.create_sheet(title=metric_name)

        # Column header row
        ws.cell(1, 1, "Dataset").font = Font(bold=True, color=C_HEADER_FG, size=9)
        ws.cell(1, 1).fill      = _xl_fill(C_HEADER_BG)
        ws.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
        ws.cell(1, 1).border    = _xl_border()
        ws.row_dimensions[1].height = 18

        for j, lbl in enumerate(col_labels, start=2):
            c = ws.cell(1, j, lbl)
            c.font      = Font(bold=True, color=C_HEADER_FG, size=9)
            c.fill      = _xl_fill(C_HEADER_BG)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border    = _xl_border()

        write_metric_block(ws, start_row=2,
                           mean_col=mean_col, sd_col=sd_col,
                           metric_label=metric_label)

        # Column widths
        ws.column_dimensions["A"].width = 14
        for j in range(2, 2 + len(METHOD_ORDER)):
            ws.column_dimensions[get_column_letter(j)].width = 13

        ws.freeze_panes = "B2"

    # ── Sheet: All metrics combined (paper-style) ────────────────────────────
    ws_all = wb.create_sheet(title="All Metrics", index=0)
    n_cols = 1 + len(METHOD_ORDER)

    # Column header
    ws_all.cell(1, 1, "Dataset").font      = Font(bold=True, color=C_HEADER_FG, size=9)
    ws_all.cell(1, 1).fill      = _xl_fill(C_HEADER_BG)
    ws_all.cell(1, 1).alignment = Alignment(horizontal="center", vertical="center")
    ws_all.cell(1, 1).border    = _xl_border()
    ws_all.row_dimensions[1].height = 18

    for j, lbl in enumerate(col_labels, start=2):
        c = ws_all.cell(1, j, lbl)
        c.font      = Font(bold=True, color=C_HEADER_FG, size=9)
        c.fill      = _xl_fill(C_HEADER_BG)
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border    = _xl_border()

    cur_row = 2
    for mean_col, sd_col, metric_label in METRICS:
        cur_row = write_metric_block(ws_all, cur_row, mean_col, sd_col, metric_label)
        cur_row += 1   # blank separator row

    ws_all.column_dimensions["A"].width = 14
    for j in range(2, 2 + len(METHOD_ORDER)):
        ws_all.column_dimensions[get_column_letter(j)].width = 13
    ws_all.freeze_panes = "B2"

    # ── Sheet: Raw data (long format) ────────────────────────────────────────
    ws_raw = wb.create_sheet(title="Raw Data")
    raw_cols = ["dataset", "method", "ACC_mean", "ACC_sd",
                "NMI_mean", "NMI_sd", "F_mean", "F_sd", "PUR_mean", "PUR_sd"]
    for j, col in enumerate(raw_cols, start=1):
        c = ws_raw.cell(1, j, col)
        c.font      = Font(bold=True, color=C_HEADER_FG, size=9)
        c.fill      = _xl_fill(C_HEADER_BG)
        c.alignment = Alignment(horizontal="center")
        c.border    = _xl_border()

    export_data = data[data["dataset"].isin(DATASET_ORDER)].copy()
    export_data["dataset"] = export_data["dataset"].map(
        lambda x: DATASET_LABELS.get(x, x))
    export_data["method"] = export_data["method"].map(
        lambda x: METHOD_LABELS.get(x, x))

    for i, (_, row_) in enumerate(export_data.iterrows(), start=2):
        bg = C_ODD_BG if i % 2 == 0 else C_EVEN_BG
        for j, col in enumerate(raw_cols, start=1):
            c = ws_raw.cell(i, j, row_.get(col, ""))
            c.fill      = _xl_fill(bg)
            c.font      = Font(size=9)
            c.alignment = Alignment(horizontal="center")
            c.border    = _xl_border()

    for j in range(1, len(raw_cols) + 1):
        ws_raw.column_dimensions[get_column_letter(j)].width = 13

    wb.save(out_path)
    print(f"  XLSX  : {out_path}")


# ════════════════════════════════════════════════════════════════════════════════
# 3. PNG table
# ════════════════════════════════════════════════════════════════════════════════
def make_png_table(data: pd.DataFrame, out_path: str) -> None:
    col_labels = [METHOD_LABELS[m] for m in METHOD_ORDER]
    datasets   = active_datasets(data)
    n_ds       = len(datasets)
    n_cols     = 1 + len(METHOD_ORDER)
    n_rows_total = 1 + len(METRICS) * (1 + n_ds)   # header + (metric_hdr + ds rows)

    fig_h = max(4, 0.36 * n_rows_total + 0.5)
    fig_w = 2.0 + 1.5 * len(METHOD_ORDER)
    fig, ax = plt.subplots(figsize=(fig_w, fig_h))
    ax.axis("off")

    HEADER_BG  = "#2C3E50"
    METRIC_BG  = "#D5E8F5"
    BEST_BG    = "#FFD580"
    ODD_BG     = "#FFFFFF"
    EVEN_BG    = "#F4F8FC"

    cell_text, cell_colors = [], []

    # Column header
    cell_text.append(["Dataset"] + col_labels)
    cell_colors.append([HEADER_BG] * n_cols)

    for mean_col, sd_col, metric_label in METRICS:
        best = best_per_row(data, mean_col)
        cell_text.append([metric_label] + [""] * len(METHOD_ORDER))
        cell_colors.append([METRIC_BG] * n_cols)

        for di, ds in enumerate(datasets):
            sub = data[data["dataset"] == ds]
            row_t  = [DATASET_LABELS[ds]]
            bg     = EVEN_BG if di % 2 == 0 else ODD_BG
            row_bg = [bg]
            for m in METHOD_ORDER:
                r_ = sub[sub["method"] == m]
                row_t.append(fmt(float(r_[mean_col].iloc[0]),
                                 float(r_[sd_col].iloc[0])) if not r_.empty else "—")
                row_bg.append(BEST_BG if best.get(ds) == m else bg)
            cell_text.append(row_t)
            cell_colors.append(row_bg)

    tbl = ax.table(cellText=cell_text, cellLoc="center",
                   loc="center", cellColours=cell_colors)
    tbl.auto_set_font_size(False)
    tbl.set_fontsize(7.5)
    tbl.scale(1, 1.35)

    for j in range(n_cols):
        tbl[0, j].set_text_props(color="white", fontweight="bold")

    row_idx = 1
    for mean_col, _, _ in METRICS:
        best = best_per_row(data, mean_col)
        tbl[row_idx, 0].set_text_props(fontweight="bold", color="#1A3A5C")
        row_idx += 1
        for ds in datasets:
            tbl[row_idx, 0].set_text_props(fontstyle="italic")
            b = best.get(ds)
            if b and b in METHOD_ORDER:
                tbl[row_idx, 1 + METHOD_ORDER.index(b)].set_text_props(
                    fontweight="bold", color="#7B3F00")
            row_idx += 1

    legend_elements = [
        mpatches.Patch(facecolor=BEST_BG, label="Best per row"),
    ]
    ax.legend(handles=legend_elements, loc="lower right", fontsize=7,
              framealpha=0.7, bbox_to_anchor=(1.0, -0.02))
    fig.suptitle("GMM-Incomplete — Clustering Performance\n"
                 "(mean ± std, missing ratios 10%–70%)",
                 fontsize=9, fontweight="bold", y=1.01)
    plt.tight_layout()
    plt.savefig(out_path, dpi=180, bbox_inches="tight")
    plt.close()
    print(f"  PNG   : {out_path}")


# ════════════════════════════════════════════════════════════════════════════════
# 4. Heatmap (all 4 metrics)
# ════════════════════════════════════════════════════════════════════════════════
def make_heatmap(data: pd.DataFrame, out_path: str) -> None:
    datasets  = active_datasets(data)
    n_ds      = len(datasets)
    n_meth    = len(METHOD_ORDER)
    n_metrics = len(METRICS)

    # Soft RdYlGn: truncate to [0.12, 0.88] to avoid overly dark red/green
    base_cmap = plt.get_cmap("RdYlGn")
    soft_colors = base_cmap(np.linspace(0.12, 0.88, 256))
    CMAP = LinearSegmentedColormap.from_list("RdYlGn_soft", soft_colors)

    # Extra bottom margin for the single shared colorbar
    fig_h = 0.72 * n_ds + 2.8
    fig_w = 3.6 * n_metrics
    fig, axes = plt.subplots(1, n_metrics, figsize=(fig_w, fig_h))
    fig.subplots_adjust(left=0.07, right=0.98, top=0.88,
                        bottom=0.28, wspace=0.38)

    last_im = None
    for ax, (mean_col, _, label) in zip(axes, METRICS):
        matrix = np.full((n_ds, n_meth), np.nan)
        for i, ds in enumerate(datasets):
            sub = data[data["dataset"] == ds]
            for j, m in enumerate(METHOD_ORDER):
                r_ = sub[sub["method"] == m]
                if not r_.empty:
                    matrix[i, j] = float(r_[mean_col].iloc[0])

        # Per-row normalisation (0 = worst in row, 1 = best in row)
        row_norm = np.full_like(matrix, np.nan)
        for i in range(n_ds):
            row = matrix[i, :]
            lo, hi = np.nanmin(row), np.nanmax(row)
            row_norm[i, :] = (row - lo) / (hi - lo) if hi > lo else np.where(np.isnan(row), np.nan, 0.5)

        last_im = ax.imshow(row_norm, aspect="auto", cmap=CMAP, vmin=0, vmax=1)

        # Cell text: black on yellow mid-range, white on strong red/green
        for i in range(n_ds):
            best_j = int(np.nanargmax(matrix[i, :]))
            for j in range(n_meth):
                if np.isnan(matrix[i, j]):
                    continue
                norm_v    = row_norm[i, j]
                txt_color = "white" if (norm_v < 0.28 or norm_v > 0.75) else "black"
                ax.text(j, i, f"{matrix[i, j]:.1f}",
                        ha="center", va="center", fontsize=8,
                        fontweight="bold" if j == best_j else "normal",
                        color=txt_color)

        ax.set_xticks(range(n_meth))
        ax.set_xticklabels([METHOD_LABELS[m] for m in METHOD_ORDER],
                           rotation=38, ha="right", fontsize=8.5)
        ax.set_yticks(range(n_ds))
        ax.set_yticklabels([DATASET_LABELS[ds] for ds in datasets], fontsize=9)
        ax.set_title(label, fontsize=11, fontweight="bold", pad=8)
        ax.tick_params(length=0)

        # Black border on best cell per row
        for i in range(n_ds):
            best_j = int(np.nanargmax(matrix[i, :]))
            ax.add_patch(plt.Rectangle((best_j - 0.5, i - 0.5), 1, 1,
                         fill=False, edgecolor="black", linewidth=1.2))

    # Single shared colorbar at the bottom (all subplots share 0–1 scale)
    norm = plt.Normalize(vmin=0, vmax=1)
    sm   = plt.cm.ScalarMappable(cmap=CMAP, norm=norm)
    sm.set_array([])
    cbar_ax = fig.add_axes([0.20, 0.10, 0.60, 0.025])
    cbar = fig.colorbar(sm, cax=cbar_ax, orientation="horizontal")
    cbar.set_ticks([0, 0.5, 1])
    cbar.set_ticklabels(["Worst in row", "Mid", "Best in row"], fontsize=9)
    cbar.ax.tick_params(length=0)
    cbar.set_label("Relative performance within each dataset row", fontsize=9, labelpad=6)

    fig.suptitle("GMM-Incomplete — Performance Heatmap  (values = mean %, averaged over missing ratios 10%–70%)\n",
                 fontsize=10.5, fontweight="bold", y=0.97)
    plt.savefig(out_path, dpi=180, bbox_inches="tight")
    plt.close()
    print(f"  Heatmap : {out_path}")


# ════════════════════════════════════════════════════════════════════════════════
# 5. Delta heatmap — Proposed vs baselines (ACC)
# ════════════════════════════════════════════════════════════════════════════════
def make_delta_heatmap(data: pd.DataFrame, out_path: str) -> None:
    datasets  = active_datasets(data)
    baselines = [m for m in METHOD_ORDER if m != "Proposed"]
    matrix    = np.full((len(datasets), len(baselines)), np.nan)

    for i, ds in enumerate(datasets):
        sub  = data[data["dataset"] == ds]
        prop = sub[sub["method"] == "Proposed"]["ACC_mean"]
        if prop.empty:
            continue
        pv = float(prop.iloc[0])
        for j, m in enumerate(baselines):
            r_ = sub[sub["method"] == m]["ACC_mean"]
            if not r_.empty:
                matrix[i, j] = pv - float(r_.iloc[0])

    vabs = np.nanmax(np.abs(matrix))
    cmap = LinearSegmentedColormap.from_list(
        "div", ["#D32F2F", "#FFFFFF", "#1A7644"], N=256)

    fig, ax = plt.subplots(figsize=(8, 0.7 * len(datasets) + 1.5))
    im = ax.imshow(matrix, aspect="auto", cmap=cmap, vmin=-vabs, vmax=vabs)

    for i in range(len(datasets)):
        for j in range(len(baselines)):
            if not np.isnan(matrix[i, j]):
                sign  = "+" if matrix[i, j] >= 0 else ""
                color = "black" if abs(matrix[i, j]) < vabs * 0.6 else "white"
                ax.text(j, i, f"{sign}{matrix[i, j]:.1f}",
                        ha="center", va="center", fontsize=8.5,
                        fontweight="bold", color=color)

    ax.set_xticks(range(len(baselines)))
    ax.set_xticklabels([METHOD_LABELS[m] for m in baselines],
                       rotation=30, ha="right", fontsize=9)
    ax.set_yticks(range(len(datasets)))
    ax.set_yticklabels([DATASET_LABELS[ds] for ds in datasets], fontsize=9)
    ax.set_title("Proposed (Ours) − Baseline  [Δ ACC %]\n"
                 "Green = Proposed better  |  Red = Proposed worse",
                 fontsize=10, fontweight="bold")
    plt.colorbar(im, ax=ax, shrink=0.8).set_label("Δ ACC (%)", fontsize=8)
    plt.tight_layout()
    plt.savefig(out_path, dpi=180, bbox_inches="tight")
    plt.close()
    print(f"  Delta   : {out_path}")


# ════════════════════════════════════════════════════════════════════════════════
# Main
# ════════════════════════════════════════════════════════════════════════════════
if __name__ == "__main__":
    if not os.path.exists(XLSX_IN):
        sys.exit(f"ERROR: input file not found:\n  {XLSX_IN}")

    data = load_data(XLSX_IN)
    print(f"Loaded {len(data)} rows | datasets: {sorted(data['dataset'].unique())}\n")

    make_latex(        data, os.path.join(OUT_DIR, "table_comparison.tex"))
    make_xlsx(         data, os.path.join(OUT_DIR, "table_comparison.xlsx"))
    make_png_table(    data, os.path.join(OUT_DIR, "table_comparison.png"))
    make_heatmap(      data, os.path.join(OUT_DIR, "heatmap_comparison.png"))
    make_delta_heatmap(data, os.path.join(OUT_DIR, "heatmap_delta_acc.png"))

    print("\nAll outputs saved to results/")
